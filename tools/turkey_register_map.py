"""
Writes the Modbus register map for the DT Turkey / CCI platform:

  docs/TURKEY_MODBUS.md                          the map as a document
  docs/turkey/Solar_Tag_List_Rev.1.2.xlsx        "PLC Address" and connection columns filled in
  docs/turkey/Solar_Inverter_Tag_List_Rev.1.2.xlsx
  docs/turkey/Solar_Modbus_Register_Map_Rev.1.0.xlsx   the map as a workbook (Connection + Tags)

Run from the repository root:  python tools/turkey_register_map.py [--host <plant-pc-vpn-ip>]
The addresses come from plantcollector.publish.build_map, so they are the same ones the
running server uses. Run it again whenever profiles/tag_catalog.yaml changes.
"""

from __future__ import annotations

import argparse
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook  # noqa: E402

from plantcollector.publish import HEADER, MAP_VERSION, QUALITY_START, VALUES_START, build_map  # noqa: E402
from plantcollector.registry import load_registry  # noqa: E402
from plantcollector.views import load_tag_catalog  # noqa: E402

WORKBOOKS = ["docs/turkey/Solar_Tag_List_Rev.1.2.xlsx", "docs/turkey/Solar_Inverter_Tag_List_Rev.1.2.xlsx"]
MAP_WORKBOOK = "docs/turkey/Solar_Modbus_Register_Map_Rev.1.0.xlsx"


def write_map_workbook(path: str, tags: list[dict], host: str, port: int, unit: int, reg=None) -> None:
    """The register map as a workbook: a Connection sheet and the Tags sheet with a readiness column."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Connection"
    rows = [
        ("Erbil solar plant - Modbus TCP register map for the CCI monitoring platform", None, None), (None, None, None),
        ("Server", "Plant data collector on the plant PC (read-only Modbus TCP server)", None),
        ("IP address", host if not host.startswith("<") else "to be provided once the connection method (VPN) is agreed", None),
        ("Port", str(port), None), ("Unit ID", str(unit), None),
        ("Function codes", "03 Read Holding Registers (04 returns the same registers). Writes are refused.", None),
        ("Addressing", "0-based protocol addresses (address 100 = holding register 40101 in 1-based notation)", None),
        ("Float tags", "float32, 2 registers, high word first (ABCD), engineering units, scale 1", None),
        ("Boolean tags", "uint16, 1 register, 0 / 1", None),
        ("Quality registers", f"one uint16 per tag at {QUALITY_START} + index: 1 good, 2 stale, 0 bad / not available", None),
        ("Data refresh", "values are read from the site devices once per hour (30 min at the fastest); please poll every 30-60 minutes", None),
        ("Register map version", str(MAP_VERSION), None), (None, None, None),
        ("Header registers", None, None),
        (f"{HEADER['last_read_time']}-{HEADER['last_read_time'] + 1}", "uint32 (high word first)", "Unix time (UTC seconds) of the last site read"),
        (str(HEADER["minutes_since_read"]), "uint16", "Minutes since the last site read (65535 = never)"),
        (str(HEADER["good_tags"]), "uint16", "Number of tags with quality good"),
        (str(HEADER["total_tags"]), "uint16", f"Total number of tags ({len(tags)})"),
        (str(HEADER["map_version"]), "uint16", "Register map version"),
        (str(HEADER["heartbeat"]), "uint16", "Heartbeat, +1 at every site read"),
        (None, None, None),
        ("Status column (Tags sheet)", "Ready = the source signal is already wired; Pending commissioning = the address is fixed, the value stays at quality 0 until commissioning of that signal is complete", None),
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 70; ws.column_dimensions["C"].width = 50
    ts = wb.create_sheet("Tags")
    ts.append(["#", "Tag name", "Address", "Registers", "Data type", "Unit", "Quality address", "Source", "Status"])
    for c in ts[1]:
        c.font = Font(bold=True)
    pending = PatternFill("solid", fgColor="FFF2CC")
    for t in tags:
        st = source_status(reg, t)
        ts.append([t["index"], t["name"], t["address"], t["registers"], "Float32 (ABCD)" if t["type"] == "float32" else "Boolean (uint16 0/1)",
                   t["unit"], t["quality_address"], f"{t['device']}.{t['point']}", st])
        if st != "Ready":
            ts.cell(row=ts.max_row, column=9).fill = pending
    for i, w in enumerate((6, 44, 10, 10, 20, 8, 15, 30, 62), start=1):
        ts.column_dimensions[get_column_letter(i)].width = w
    ts.freeze_panes = "A2"
    wb.save(path)


def write_markdown(tags: list[dict], host: str, port: int, unit: int, reg=None) -> str:
    lines = [
        "# Modbus TCP register map for the CCI / DT Turkey platform",
        "",
        f"Server: the plant data collector on the plant PC — `{host}:{port}`, unit id `{unit}`, Modbus TCP, "
        "**read-only** (FC03 Read Holding Registers; FC04 returns the same registers; every write is refused with an exception).",
        "",
        "The values are the collector's last read of the site devices: **once per hour** (30 minutes at the fastest). "
        "Please poll every 30–60 minutes; faster polling only returns the same values again. Registers 0–2 give the time "
        "of the last site read, so a poll a few minutes after the full hour always gets fresh data. Polling never reaches "
        "the SAU or the InteliNeo.",
        "",
        f"Register map version {MAP_VERSION}. Addresses are 0-based protocol addresses "
        "(address 100 = holding register 40101 in 1-based notation).",
        "",
        "## Header",
        "",
        "| Address | Type | Meaning |",
        "|---|---|---|",
        f"| {HEADER['last_read_time']}–{HEADER['last_read_time'] + 1} | uint32 (high word first) | Unix time (UTC seconds) of the last site read |",
        f"| {HEADER['minutes_since_read']} | uint16 | Minutes since the last site read (65535 = never) |",
        f"| {HEADER['good_tags']} | uint16 | Number of tags with quality *good* |",
        f"| {HEADER['total_tags']} | uint16 | Total number of tags ({len(tags)}) |",
        f"| {HEADER['map_version']} | uint16 | Register map version |",
        f"| {HEADER['heartbeat']} | uint16 | Heartbeat, +1 at every site read |",
        "",
        "## Encoding",
        "",
        "- Float tags: **float32**, 2 registers, high word first (ABCD), engineering units as in the Unit column, scale 1.",
        "- Boolean tags: **uint16**, 1 register, 0 or 1.",
        f"- Quality: one uint16 per tag at `{QUALITY_START} + index`: **1** good, **2** stale (older than two read intervals), "
        "**0** bad / device not answering / no data yet. While a tag is bad its value registers hold the last good value "
        "(configurable to zero or NaN).",
        "",
        "## Tags",
        "",
        "| # | Tag name | Address | Registers | Type | Unit | Quality address | Source | Status |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for t in tags:
        lines.append(f"| {t['index']} | {t['name']} | {t['address']} | {t['registers']} | {t['type']} | {t['unit']} | "
                     f"{t['quality_address']} | {t['device']}.{t['point']} | {source_status(reg, t)} |")
    lines += ["", f"Value block: {VALUES_START}–{tags[-1]['address'] + tags[-1]['registers'] - 1}. "
                  f"Quality block: {QUALITY_START}–{QUALITY_START + len(tags) - 1}.", ""]
    return "\n".join(lines)


def source_status(reg, t: dict) -> str:
    """Whether the tag's source point is wired to a register on the plant yet."""
    d = reg.devices.get(t["device"]) if reg else None
    p = d.point(t["point"]) if d else None
    if p is not None and ("address" in p.raw or (p.raw.get("derived") and any("address" in q.raw for q in d.points))):
        return "Ready"
    return "Pending commissioning"


def fill_workbook(path: str, by_name: dict[str, dict], host: str, port: int, unit: int, reg=None) -> int:
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h.lower().replace("_", " "): i + 1 for i, h in enumerate(head)}

    def put(row: int, name: str, value):
        for key, idx in col.items():
            if key == name:
                ws.cell(row=row, column=idx, value=value)
                return

    n = 0
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=col.get("tag name", 1)).value
        t = by_name.get(str(name).strip()) if name else None
        if not t:
            continue
        n += 1
        put(row, "plc address", t["address"])
        put(row, "source device", "Plant PC data collector (Modbus TCP server, read-only)")
        put(row, "protocol", "Modbus TCP")
        put(row, "ip:port", f"{host}:{port}")
        put(row, "unit id", unit)
        put(row, "register / fc", f"Holding (FC03), {t['registers']} reg" + ("s" if t["registers"] > 1 else ""))
        put(row, "register type", f"Holding (FC03), {t['registers']} reg" + ("s" if t["registers"] > 1 else ""))
        put(row, "data type", "Float32 (ABCD)" if t["type"] == "float32" else "Boolean (uint16 0/1)")
        put(row, "scale", 1)
        put(row, "status", f"{source_status(reg, t)} - quality register {t['quality_address']}")
    wb.save(path)
    return n


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--host", default="<plant-pc-vpn-ip>")
    ap.add_argument("--port", type=int, default=502)
    ap.add_argument("--unit", type=int, default=1)
    ap.add_argument("--config", default="config/site.plant.yaml", help="plant config, for the per-tag readiness column")
    a = ap.parse_args()
    reg = load_registry(os.path.join(ROOT, a.config)) if os.path.exists(os.path.join(ROOT, a.config)) else None
    tags = build_map(load_tag_catalog())
    by_name = {t["name"]: t for t in tags}
    md = os.path.join(ROOT, "docs", "TURKEY_MODBUS.md")
    with open(md, "w", encoding="utf-8") as f:
        f.write(write_markdown(tags, a.host, a.port, a.unit, reg))
    print(f"wrote {md} ({len(tags)} tags)")
    for wb in WORKBOOKS:
        n = fill_workbook(os.path.join(ROOT, wb), by_name, a.host, a.port, a.unit, reg)
        print(f"filled {n} rows in {wb}")
    write_map_workbook(os.path.join(ROOT, MAP_WORKBOOK), tags, a.host, a.port, a.unit, reg)
    print(f"wrote {MAP_WORKBOOK}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
