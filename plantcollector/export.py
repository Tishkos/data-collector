"""
Exports: a wide time table (one column per device.point) as CSV or Excel,
the daily report workbook, and the Turkey tag list with live values.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import Collector
from .views import catalog_values, local_tz, series, series_specs

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=13)
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="666666")


def _style_header(ws, row: int = 1):
    for cell in ws[row]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autowidth(ws, minimum: int = 10, maximum: int = 40):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimum, min(maximum, width + 2))


# --- wide table ---------------------------------------------------------------------------------

def wide_table(c: Collector, since: datetime, until: datetime, step: int) -> tuple[list[str], list[list]]:
    reg, st = c.registry, c.store
    present = set(st.pairs_in_range(since, until))
    cols: list[tuple[str, str, str]] = []           # (device, point, unit)
    for d in reg.devices.values():
        for p in d.points:
            if (d.id, p.id) in present:
                cols.append((d.id, p.id, p.unit))
    data = st.resample(since, until, step, [(d, p) for d, p, _ in cols])
    tz = local_tz(reg)
    b0 = int(since.timestamp()) // step * step
    b1 = int(until.timestamp()) // step * step
    headers = ["time_utc", f"time_local ({reg.site.timezone})"] + [f"{d}.{p}" + (f" [{u}]" if u else "") for d, p, u in cols]
    rows: list[list] = []
    for b in range(b0, b1 + step, step):
        t = datetime.fromtimestamp(b, tz=timezone.utc)
        row = [t.strftime("%Y-%m-%d %H:%M:%S"), t.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S")]
        any_val = False
        for d, p, _u in cols:
            v = data.get((d, p), {}).get(b)
            if v is not None:
                any_val = True
                v = round(v, 4)
            row.append(v)
        if any_val:
            rows.append(row)
    return headers, rows


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(headers)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(c: Collector, headers: list[str], rows: list[list], since: datetime, until: datetime, step: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _style_header(ws)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    _autowidth(ws, 12, 28)

    # summary sheet for the exported window
    s = wb.create_sheet("Summary")
    reg = c.registry
    s["A1"], s["A1"].font = f"{reg.site.name} — data export", TITLE_FONT
    s["A2"], s["A2"].font = f"window {since.strftime('%Y-%m-%d %H:%M')} to {until.strftime('%Y-%m-%d %H:%M')} UTC, {step} s averages", NOTE_FONT
    s["A3"], s["A3"].font = "values are read-only polls of the plant devices; quality 'good' samples only", NOTE_FONT
    s.append([])
    s.append(["quantity", "value", "unit"])
    _style_header(s, 5)
    ser = series(c, since, until, step)
    hours = step / 3600.0
    def kwh(key):
        v = ser["series"].get(key)
        return round(sum(x for x in v if x is not None) * hours, 1) if v else None
    load, pv, gen, bess = kwh("load_kw"), kwh("pv_kw"), kwh("genset_kw"), ser["series"].get("bess_kw")
    lines = [("site load energy", load, "kWh"), ("solar energy", pv, "kWh"), ("grid energy", kwh("mains_kw"), "kWh"), ("genset energy", gen, "kWh")]
    if bess:
        lines.append(("battery discharged", round(sum(x for x in bess if x and x > 0) * hours, 1), "kWh"))
        lines.append(("battery charged", round(-sum(x for x in bess if x and x < 0) * hours, 1), "kWh"))
    if load and pv is not None:
        lines.append(("solar share of load", round(100 * pv / load, 1), "%"))
    for g in reg.by_kind("genset"):
        f = ser["series"].get(f"{g.id}/fuel_rate_lph")
        if f:
            lines.append((f"{g.id} fuel", round(sum(x for x in f if x is not None) * hours, 1), "L"))
    for q, v, u in lines:
        s.append([q, v, u])
    for row in s.iter_rows(min_row=6):
        for cell in row:
            cell.font = BODY_FONT
    _autowidth(s, 14, 40)

    # point dictionary
    p = wb.create_sheet("Points")
    p.append(["column", "device", "kind", "point", "unit", "type", "description", "modbus address", "register"])
    for d in reg.devices.values():
        for pt in d.points:
            p.append([f"{d.id}.{pt.id}", d.id, d.kind, pt.id, pt.unit, pt.type, pt.description,
                      pt.raw.get("address"), pt.raw.get("kind")])
    _style_header(p)
    for row in p.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    _autowidth(p, 10, 48)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- daily report -----------------------------------------------------------------------------------

def daily_report_xlsx(c: Collector, day: datetime | None = None) -> bytes:
    rep = c.daily_totals(day)
    reg = c.registry
    tz = local_tz(reg)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = f"{reg.site.name} — daily report {rep['date']}", TITLE_FONT
    ws["A2"], ws["A2"].font = f"day boundaries in UTC; local time zone {reg.site.timezone}", NOTE_FONT
    ws.append([])
    ws.append(["quantity", "value", "unit"])
    _style_header(ws, 4)
    t = rep["totals"]
    for q, k, u in [("site load energy", "load_kwh", "kWh"), ("solar energy", "pv_kwh", "kWh"),
                    ("solar available", "pv_available_kwh", "kWh"), ("solar curtailed", "pv_curtailed_kwh", "kWh"),
                    ("solar share of load", "pv_share_pct", "%"), ("genset energy", "genset_kwh", "kWh"),
                    ("fuel used", "fuel_l", "L"), ("fuel cost", "fuel_cost", "currency"),
                    ("fuel saved vs no solar (estimate)", "fuel_saved_vs_no_solar_l", "L"),
                    ("grid energy (import)", "grid_kwh", "kWh"), ("grid available", "grid_hours", "h"),
                    ("grid share of load", "grid_share_pct", "%"),
                    ("battery discharged", "bess_discharged_kwh", "kWh"), ("battery charged", "bess_charged_kwh", "kWh")]:
        ws.append([q, t.get(k), u])
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.font = BODY_FONT
    _autowidth(ws, 14, 40)

    g = wb.create_sheet("Gensets")
    g.append(["genset", "energy kWh", "fuel L", "run hours", "avg load kW", "L per kWh"])
    for gid, v in rep["gensets"].items():
        g.append([gid, v["energy_kwh"], v["fuel_l"], v["run_hours"], v["avg_load_kw"], v["l_per_kwh"]])
    _style_header(g)
    _autowidth(g)

    # hourly energy table
    start = datetime.fromisoformat(rep["date"]).replace(tzinfo=timezone.utc)
    ser = series(c, start, start + timedelta(days=1), 3600)
    h = wb.create_sheet("Hourly")
    keys = [k for k in ("load_kw", "pv_kw", "genset_kw", "bess_kw", "mains_kw") if k in ser["series"]]
    h.append(["hour (UTC)", "hour (local)"] + [f"{k.replace('_kw', '')} kWh" for k in keys])
    for i, ms in enumerate(ser["t"][:-1]):
        tt = datetime.fromtimestamp(ms / 1000, tz=timezone.utc)
        h.append([tt.strftime("%H:%M"), tt.astimezone(tz).strftime("%H:%M")] +
                 [None if ser["series"][k][i] is None else round(ser["series"][k][i], 1) for k in keys])
    _style_header(h)
    _autowidth(h)

    e = wb.create_sheet("Events")
    e.append(["time (UTC)", "severity", "device", "code", "message"])
    for ev in rep["events"]:
        e.append([ev["ts"][:19].replace("T", " "), ev["severity"], ev["device_id"], ev["code"], ev["message"]])
    _style_header(e)
    _autowidth(e, 10, 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- Turkey tag list with values -----------------------------------------------------------------------

def tags_xlsx(c: Collector, catalog: list[dict]) -> bytes:
    rows = catalog_values(c, catalog)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tags"
    ws.append(["tag name", "type", "unit", "value", "quality", "timestamp (UTC)", "source device", "source point", "mapped"])
    for r in rows:
        ws.append([r["name"], r["type"], r["unit"], r["value"], r["quality"], (r["ts"] or "")[:19].replace("T", " "),
                   r["device"], r["point"], "yes" if r["mapped"] else "no"])
    _style_header(ws)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    _autowidth(ws, 10, 40)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def tags_csv(c: Collector, catalog: list[dict]) -> bytes:
    rows = catalog_values(c, catalog)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["tag name", "type", "unit", "value", "quality", "timestamp_utc", "device", "point"])
    for r in rows:
        w.writerow([r["name"], r["type"], r["unit"], r["value"], r["quality"], r["ts"], r["device"], r["point"]])
    return buf.getvalue().encode("utf-8-sig")


__all__ = ["wide_table", "to_csv", "to_xlsx", "daily_report_xlsx", "tags_xlsx", "tags_csv", "series_specs"]
