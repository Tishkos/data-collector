import io
import os
import tempfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture
def client(monkeypatch):
    tmp = tempfile.mkdtemp()
    monkeypatch.setenv("QS_SITE_CONFIG", os.path.join(ROOT, "config", "site.example.yaml"))
    monkeypatch.setenv("QS_DATA_DIR", tmp)
    from qs_hybrid.api import app as mod
    mod._site = None
    mod._autopilot = None
    c = TestClient(mod.app)
    mod.site().advance(2 * 3600, step=300)      # 2 h of simulated history
    return c


def test_dashboard_page_and_static(client):
    r = client.get("/")
    assert r.status_code == 200 and "Plant dashboard" in r.text
    assert client.get("/static/chart.umd.js").status_code == 200


def test_summary_series_tags(client):
    s = client.get("/api/v1/local/summary").json()
    assert s["kpi"]["load_kw"] > 0 and s["kpi"]["gensets_total"] == 2
    assert any(d["id"] == "pv-01" and d["comms"] == "ok" for d in s["devices"])
    r = client.get("/api/v1/local/series?hours=3").json()
    assert len(r["series"]["load_kw"]) >= 10 and len(r["series"]["soc_pct"]) >= 10
    t = client.get("/api/v1/local/tags").json()
    assert any(x["device"] == "bess-01" and x["point"] == "soc_pct" and x["quality"] == "good" for x in t)


def test_csv_and_xlsx_export(client):
    r = client.get("/api/v1/local/export.csv?step=300")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.csv"')
    lines = r.text.strip().splitlines()
    assert lines[0].startswith("timestamp,") and "bess-01/soc_pct [%]" in lines[0] and len(lines) > 10

    r = client.get("/api/v1/local/export.xlsx?step=300")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Data", "Summary"]
    ws = wb["Data"]
    assert ws["A1"].value == "timestamp" and ws.max_row > 10 and ws.max_column > 20
    assert ws["A1"].font.name == "Arial"
    sm = wb["Summary"]
    assert sm["A1"].value == "Site"

    r = client.get("/api/v1/local/report.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.active["A1"].value == "Site"


def test_export_range_validation(client):
    r = client.get("/api/v1/local/export.csv?from=2026-10-13T10:00:00&to=2026-10-13T09:00:00")
    assert r.status_code == 400
