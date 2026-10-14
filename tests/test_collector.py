"""End-to-end checks on the demo plant: store, views, exports, API. No network."""

import io
import os
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from plantcollector.api import create_app
from plantcollector.core import Collector
from plantcollector import export, views

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = os.path.join(ROOT, "config", "site.yaml")


@pytest.fixture(scope="module")
def collector(tmp_path_factory):
    return Collector(CFG, str(tmp_path_factory.mktemp("data")), seed_hours=30, fresh=True)


@pytest.fixture(scope="module")
def client(collector):
    return TestClient(create_app(collector))


def test_registry_is_read_only(collector):
    for d in collector.registry.devices.values():
        assert not hasattr(d, "commands")
    import plantcollector.adapters.modbus_tcp as m
    src = open(m.__file__).read()
    assert "write_register" not in src and "write_coil" not in src


def test_seeded_history_and_totals(collector):
    assert collector.store.count() > 1000
    t = collector.daily_totals()["totals"]
    assert t["load_kwh"] > 0 and t["pv_kwh"] > 0
    assert abs(t["load_kwh"] - (t["pv_kwh"] + t["genset_kwh"] + t["grid_kwh"])) / t["load_kwh"] < 0.08


def test_summary_and_series(collector):
    s = views.summary(collector)
    assert s["mode"] == "read-only" and s["simulated"] is True
    assert s["kpi"]["gensets_total"] == 5
    assert all(d["comms"] == "ok" for d in s["devices"])
    ser = views.series(collector, collector.now() - timedelta(hours=6), collector.now())
    assert ser["step"] >= 30 and len(ser["t"]) == len(ser["series"]["load_kw"])
    assert any(v is not None for v in ser["series"]["mains_kw"])
    assert "genset-01/power_kw" in ser["series"]


def test_events_logged(collector):
    codes = {e["code"] for e in collector.store.events(collector.now() - timedelta(hours=30))}
    assert {"COMMS_LOST", "COMMS_OK", "ALARM", "ALARM_OK", "GRID_OFF", "GRID_OK"} <= codes


def test_wide_export_csv_and_xlsx(collector):
    since, until = collector.now() - timedelta(hours=3), collector.now()
    headers, rows = export.wide_table(collector, since, until, 300)
    assert headers[0] == "time_utc" and "meter-01.load_kw [kW]" in headers
    assert 30 <= len(rows) <= 38
    csv = export.to_csv(headers, rows).decode("utf-8-sig").splitlines()
    assert csv[0].startswith("time_utc,") and len(csv) == len(rows) + 1
    wb = load_workbook(io.BytesIO(export.to_xlsx(collector, headers, rows, since, until, 300)))
    assert wb.sheetnames == ["Data", "Summary", "Points"]
    assert wb["Data"].max_row == len(rows) + 1


def test_daily_report(collector):
    wb = load_workbook(io.BytesIO(export.daily_report_xlsx(collector)))
    assert wb.sheetnames == ["Summary", "Gensets", "Hourly", "Events"]
    assert wb["Hourly"].max_row == 25


def test_turkey_tags(collector):
    cat = views.load_tag_catalog()
    assert len(cat) == 305
    vals = views.catalog_values(collector, cat)
    assert vals[0]["name"].startswith("INVERTER_1.")
    wb = load_workbook(io.BytesIO(export.tags_xlsx(collector, cat)))
    assert wb["Tags"].max_row == 306


def test_api(client):
    assert client.get("/").status_code == 200
    assert client.get("/api/v1/health").json()["ok"] is True
    s = client.get("/api/v1/local/summary").json()
    assert s["kpi"]["load_kw"] is not None
    r = client.get("/api/v1/local/series?hours=6&keys=load_kw,pv_kw").json()
    assert set(r["series"]) == {"load_kw", "pv_kw"}
    assert client.get("/api/v1/local/series?hours=99999").status_code == 422
    a = client.get("/api/v1/local/alarms?hours=48").json()
    assert "events" in a and "active" in a
    x = client.get("/api/v1/local/export.xlsx?hours=2&step=300")
    assert x.status_code == 200 and x.headers["content-disposition"].endswith('.xlsx"')
    c = client.get("/api/v1/local/export.csv?hours=2&step=300")
    assert c.status_code == 200 and c.text.startswith("﻿time_utc,")
    assert client.get("/api/v1/local/report.xlsx").status_code == 200
    t = client.get("/api/v1/tags").json()
    assert t["count"] == 305
    assert client.get("/api/v1/tags/INVERTER_1.INV1_AC_POWER").json()["device"] == "inv-01"
    assert client.get("/api/v1/tags/NOPE").status_code == 404
    assert client.get("/api/v1/tags.csv").status_code == 200
    assert len(client.get("/api/v1/local/tags").json()) == sum(len(d.points) for d in client.app.state.collector.registry.devices.values())


def test_no_write_endpoints(client):
    for path in ("/api/v1/local/summary", "/api/v1/tags"):
        assert client.post(path).status_code == 405


def test_timezone_fallback_without_tz_database(monkeypatch):
    """Windows has no tz database unless tzdata is installed; exports must still work."""
    import zoneinfo

    def broken(*a, **k):
        raise zoneinfo.ZoneInfoNotFoundError("no tz database")
    monkeypatch.setattr(views, "ZoneInfo", broken)
    from plantcollector.registry import Registry, SiteConfig
    tz = views.local_tz(Registry(site=SiteConfig(site_id="x", timezone="Asia/Baghdad")))
    assert tz.utcoffset(None).total_seconds() == 3 * 3600
    assert views.local_tz(Registry(site=SiteConfig(site_id="x", timezone="Nowhere/Town"))).utcoffset(None).total_seconds() == 0


def test_settings_endpoint_demo_is_locked(client):
    st = client.get("/api/v1/local/settings").json()
    assert [o["value"] for o in st["options"]] == [1800, 3600, 7200] and st["min_s"] == 1800
    assert st["simulated"] is True
    assert client.post("/api/v1/local/settings", json={"poll_interval_s": 3600}).status_code == 400


def test_settings_persist_for_real_config(tmp_path):
    """A non-simulated config: the dashboard setting is validated, applied at once and reloaded at start."""
    cfg = tmp_path / "site.yaml"
    cfg.write_text("site: {site_id: t, poll_interval_s: 3600}\ndevices:\n"
                   "  - {id: m, profile: meter_site.yaml, adapter: modbus_tcp, connection: {host: 127.0.0.1, port: 15999, timeout: 1}}\n")
    c = Collector(str(cfg), str(tmp_path), initial_poll=False)
    client = TestClient(create_app(c))
    assert client.get("/api/v1/local/settings").json()["source"] == "config"
    for bad in (60, 900, 1799, "x", None):
        assert client.post("/api/v1/local/settings", json={"poll_interval_s": bad}).status_code == 400
    r = client.post("/api/v1/local/settings", json={"poll_interval_s": 1800})
    assert r.status_code == 200 and r.json()["poll_interval_s"] == 1800 and r.json()["source"] == "settings.json"
    assert c.poll_interval_s == 1800 and (tmp_path / "settings.json").exists()
    assert client.get("/api/v1/local/summary").json()["poll_interval_s"] == 1800
    c2 = Collector(str(cfg), str(tmp_path), initial_poll=False)
    assert c2.poll_interval_s == 1800 and c2.settings_source == "settings.json"


def test_run_loop_follows_the_setting(tmp_path):
    import threading, time
    from plantcollector.core import run_loop
    cfg = tmp_path / "site.yaml"
    cfg.write_text("site: {site_id: t, poll_interval_s: 3600}\ndevices:\n"
                   "  - {id: m, profile: meter_site.yaml, adapter: modbus_tcp, connection: {host: 127.0.0.1, port: 15999, timeout: 1}}\n")
    c = Collector(str(cfg), str(tmp_path), initial_poll=False)
    stop = threading.Event()
    t = threading.Thread(target=run_loop, args=(c, None, stop), daemon=True); t.start()
    time.sleep(2.5)
    assert c.polls == 1                      # first read at once, then nothing for an hour
    c.registry.site.poll_interval_s = 1.0    # (what set_poll_interval does, below the public minimum, for the test)
    time.sleep(3.5)
    stop.set()
    assert c.polls >= 3


def test_modbus_publisher_serves_catalog_read_only(tmp_path):
    """The Turkey-facing Modbus server: header, float32 values, quality, booleans, writes refused."""
    import struct, time
    from pymodbus.client import ModbusTcpClient
    from plantcollector.adapters import reading
    from plantcollector.publish import ModbusPublisher, QUALITY_START, VALUES_START, build_map
    from plantcollector.registry import PointDefinition

    cfg = tmp_path / "site.yaml"
    cfg.write_text("site: {site_id: t, poll_interval_s: 3600}\ndevices:\n"
                   "  - {id: inv-01, profile: inverter_sineng.yaml, adapter: modbus_tcp, connection: {host: 127.0.0.1, port: 15999, timeout: 1}}\n"
                   "  - {id: neo-01, profile: intelineo_master.yaml, adapter: modbus_tcp, connection: {host: 127.0.0.1, port: 15999, timeout: 1}}\n")
    c = Collector(str(cfg), str(tmp_path), initial_poll=False)
    pub = ModbusPublisher(c, bind="127.0.0.1", port=15031, unit_id=1)
    m = {t["name"]: t for t in pub.map}
    assert m["INVERTER_1.INV1_AC_POWER"]["address"] == VALUES_START and m["INVERTER_1.INV1_AC_POWER"]["registers"] == 2
    assert all(t["address"] < QUALITY_START for t in pub.map) and len(pub.map) == 305
    assert build_map(pub.catalog) == pub.map          # deterministic
    pub.start()
    try:
        c.store.ingest("inv-01", {"ac_power_kw": reading(PointDefinition(id="ac_power_kw", unit="kW"), 71.6)})
        c.store.ingest("neo-01", {"pvcb_closed": reading(PointDefinition(id="pvcb_closed", type="boolean"), True)})
        c.last_poll_at = datetime.now(timezone.utc)
        pub.refresh()
        cl = ModbusTcpClient("127.0.0.1", port=15031); assert cl.connect()
        hdr = cl.read_holding_registers(address=0, count=7, device_id=1).registers
        assert abs(((hdr[0] << 16) | hdr[1]) - time.time()) < 60 and hdr[2] == 0 and hdr[3] == 2 and hdr[4] == 305 and hdr[5] == 1
        t = m["INVERTER_1.INV1_AC_POWER"]
        regs = cl.read_holding_registers(address=t["address"], count=2, device_id=1).registers
        assert abs(struct.unpack(">f", struct.pack(">HH", *regs))[0] - 71.6) < 1e-4
        assert cl.read_holding_registers(address=t["quality_address"], count=1, device_id=1).registers == [1]
        b = m["Thermal_Magnetic_Switch_On_(LV PANEL)"]
        assert b["registers"] == 1 and cl.read_holding_registers(address=b["address"], count=1, device_id=1).registers == [1]
        other = m["INVERTER_2.INV2_AC_POWER"]
        assert cl.read_holding_registers(address=other["quality_address"], count=1, device_id=1).registers == [0]
        assert cl.read_input_registers(address=t["address"], count=2, device_id=1).registers == regs
        assert cl.write_register(address=t["address"], value=1, device_id=1).isError()
        assert cl.write_registers(address=t["address"], values=[1, 2], device_id=1).isError()
        assert cl.read_holding_registers(address=t["address"], count=2, device_id=1).registers == regs
        assert cl.read_holding_registers(address=0, count=2, device_id=9).isError()
        cl.close()
        assert pub.status()["reads"] >= 6 and pub.status()["rejected_writes"] == 0   # refused by the library before our code
    finally:
        pub.stop()
